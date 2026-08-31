#!/usr/bin/env python
"""Normalise a data model from various inputs into model.schema.json shape.

    python parse_model.py schema.sql            -o model.json [--dialect oracle]
    python parse_model.py columns.csv           -o model.json --format infoschema-csv
    python parse_model.py dictionary.xlsx       -o model.json --format excel-dd
    python parse_model.py models/schema.yml     -o model.json --format dbt-yaml
    python parse_model.py domain_model.html     -o model.json --format mendix-html
    python parse_model.py hand_built.json       -o model.json --format json

The governing rule for every adapter: **absent information stays absent.** If the input
does not say whether a column is nullable, `nullable` is null — never False because
False is the common case. Checks that depend on a null field stand down. An adapter that
guesses produces a model that scores well for the wrong reason, which is worse than one
that reports gaps, so every gap is recorded in `meta.parse_warnings` and surfaces in the
report as a review-input gap.

Adding an adapter: emit the schema shape, register it in ADAPTERS, add a fixture, and
run selftest.py. Validation against assets/model.schema.json is what stops a new adapter
silently dropping half the model.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SCHEMA_PATH = os.path.join(HERE, "..", "assets", "model.schema.json")


def _blank_model(source_format, name=None):
    return {
        "meta": {
            "name": name,
            "paradigm": None,
            "paradigm_evidence": None,
            "level": None,
            "dialect": None,
            "source_format": source_format,
            "source_files": [],
            "has_requirements_doc": False,
            "has_diagram": False,
            "has_data_profile": False,
            "parse_warnings": [],
        },
        "tables": [],
        "enums": [],
        "associations": [],
        "glossary": {},
    }


def _table(name, schema=None, kind="table"):
    return {"schema": schema, "name": name, "kind": kind, "role": None,
            "description": None, "grain": None, "row_count": None,
            "classification": None, "owner": None, "retention": None,
            "partitioning": None, "persistable": None, "generalization": None,
            "system_members": None, "access_rules": None, "validation_rules": None,
            "primary_key": None, "unique_constraints": [], "checks": [],
            "indexes": [], "foreign_keys": [], "columns": []}


def _column(name, **kw):
    col = {"name": name, "type": None, "length": None, "precision": None,
           "scale": None, "nullable": None, "default": None, "description": None,
           "is_pk": None, "is_identity": None, "generated": None, "expression": None,
           "calculated": None, "classification": None, "used_in": None}
    col.update({k: v for k, v in kw.items() if k in col})
    return col


_TYPE_RE = re.compile(r"^\s*([A-Za-z][A-Za-z0-9_ ]*?)\s*(?:\(\s*(\d+)\s*(?:,\s*(\d+)\s*)?\))?\s*$")


def _split_type(raw):
    """('VARCHAR2(50)') -> (length=50, precision=None, scale=None) etc."""
    if not raw:
        return None, None, None
    m = _TYPE_RE.match(str(raw))
    if not m:
        return None, None, None
    base, a, b = m.group(1).strip().lower(), m.group(2), m.group(3)
    if a is None:
        return None, None, None
    a = int(a)
    if b is not None:
        return None, a, int(b)
    if base in {"numeric", "decimal", "number", "dec", "float"}:
        return None, a, None
    return a, None, None


def _truthy(v):
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in {"y", "yes", "true", "t", "1", "x", "required", "mandatory", "not null"}:
        return True
    if s in {"n", "no", "false", "f", "0", "", "optional", "nullable", "null"}:
        return False
    return None


# ======================================================================================
# DDL
# ======================================================================================

def parse_ddl(path, dialect=None):
    model = _blank_model("ddl" + (f"-{dialect}" if dialect else ""),
                         name=os.path.basename(path))
    model["meta"]["dialect"] = dialect
    model["meta"]["level"] = "physical"
    with open(path, encoding="utf-8-sig", errors="replace") as fh:
        sql = fh.read()

    used_sqlglot = False
    try:
        import logging

        import sqlglot
        from sqlglot import exp

        # sqlglot logs "contains unsupported syntax, falling back to Command" straight to
        # stderr for statements like CREATE EXTENSION. Harmless, but it lands in the
        # middle of a review's console output where it reads like a failure. Capture it
        # into parse_warnings instead, which is where anything the parser could not fully
        # read belongs.
        class _Collect(logging.Handler):
            def __init__(self, sink):
                super().__init__()
                self.sink = sink

            def emit(self, record):
                self.sink.append(f"sqlglot: {record.getMessage()}")

        _lg = logging.getLogger("sqlglot")
        _handler = _Collect(model["meta"]["parse_warnings"])
        _lg.addHandler(_handler)
        _lg.propagate = False
        used_sqlglot = True
    except ImportError:
        model["meta"]["parse_warnings"].append(
            "sqlglot is not installed, so DDL was parsed with regular expressions. "
            "Inline constraints and unusual syntax may have been missed; install "
            "sqlglot for a full parse.")

    by_name = {}

    if used_sqlglot:
        try:
            statements = sqlglot.parse(sql, read=dialect) if dialect else sqlglot.parse(sql)
        except Exception as e:
            model["meta"]["parse_warnings"].append(
                f"sqlglot could not parse the file ({type(e).__name__}); fell back to "
                f"regular expressions.")
            used_sqlglot = False
            statements = []
        for st in statements or []:
            if not isinstance(st, exp.Create):
                continue
            kind = (st.args.get("kind") or "").upper()
            if kind not in {"TABLE", "VIEW"}:
                continue
            tbl = st.find(exp.Table)
            if tbl is None:
                continue
            tname = tbl.name
            tschema = tbl.db or None
            t = _table(tname, tschema, "view" if kind == "VIEW" else "table")
            schema_node = st.this
            exprs = schema_node.expressions if hasattr(schema_node, "expressions") else []
            for e in exprs:
                if isinstance(e, exp.ColumnDef):
                    raw = e.args.get("kind")
                    raw_s = raw.sql(dialect=dialect) if raw is not None else None
                    length, prec, scale = _split_type(raw_s)
                    col = _column(e.name, type=raw_s, length=length,
                                  precision=prec, scale=scale)
                    for c in e.constraints or []:
                        k = c.kind
                        if isinstance(k, exp.NotNullColumnConstraint):
                            col["nullable"] = bool(k.args.get("allow_null"))
                        elif isinstance(k, exp.PrimaryKeyColumnConstraint):
                            col["is_pk"] = True
                            t["primary_key"] = (t["primary_key"] or []) + [e.name]
                            col["nullable"] = False
                        elif isinstance(k, exp.UniqueColumnConstraint):
                            t["unique_constraints"].append([e.name])
                        elif isinstance(k, exp.DefaultColumnConstraint):
                            col["default"] = k.this.sql(dialect=dialect) if k.this else None
                        elif isinstance(k, exp.GeneratedAsIdentityColumnConstraint):
                            col["is_identity"] = True
                        elif isinstance(k, exp.ComputedColumnConstraint):
                            col["generated"] = True
                            col["expression"] = k.this.sql(dialect=dialect) if k.this else None
                        elif isinstance(k, exp.CommentColumnConstraint):
                            col["description"] = (k.this.name if k.this else None)
                        elif isinstance(k, exp.CheckColumnConstraint):
                            t["checks"].append({"name": None,
                                                "expression": k.this.sql(dialect=dialect)})
                    if col["nullable"] is None:
                        # Absence of NOT NULL in DDL genuinely means nullable.
                        col["nullable"] = True
                    t["columns"].append(col)
                else:
                    # A named constraint arrives wrapped: Constraint(this=Identifier,
                    # expressions=[ForeignKey(...)]). Unwrap so that named and unnamed
                    # constraints take exactly the same path — otherwise every FK
                    # written with CONSTRAINT ... FOREIGN KEY silently disappears, and
                    # a schema with full referential integrity scores as having none.
                    if isinstance(e, exp.Constraint):
                        name = e.name or None
                        inner = list(e.expressions or [])
                    else:
                        name, inner = None, [e]
                    for node in inner:
                        _table_constraint(node, t, name, dialect, exp)
            by_name[tname.lower()] = t

    if not by_name:
        _regex_tables(sql, by_name, model)

    # These statements are outside CREATE TABLE and are handled by regex regardless,
    # because sqlglot's coverage of vendor ALTER/COMMENT syntax is uneven and a missed
    # constraint would read as a real defect.
    _regex_alters(sql, by_name)
    _regex_indexes(sql, by_name)
    _regex_comments(sql, by_name)

    model["tables"] = list(by_name.values())
    model["meta"]["source_files"] = [os.path.abspath(path)]
    if not model["tables"]:
        model["meta"]["parse_warnings"].append("no CREATE TABLE statements were found")
    if used_sqlglot:
        _lg.removeHandler(_handler)
        _lg.propagate = True
    return model


def _ref_actions(fk, exp):
    """Pull ON DELETE / ON UPDATE out of a ForeignKey node.

    sqlglot has moved these between releases: older versions expose `delete` and
    `update` args, 30.x carries them as strings in `options` on the ForeignKey or on the
    nested Reference. Read all three so an upgrade does not quietly turn every
    referential action into None — which would make INT-003 fire on a schema that
    declares them properly.
    """
    on_del = fk.args.get("delete")
    on_upd = fk.args.get("update")
    opts = list(fk.args.get("options") or [])
    ref = fk.args.get("reference")
    if ref is not None:
        opts += list(ref.args.get("options") or [])
        txt = ref.sql().upper()
    else:
        txt = ""
    blob = " ".join(str(o).upper() for o in opts) + " " + txt
    if not on_del:
        m = re.search(r"ON DELETE (CASCADE|SET NULL|SET DEFAULT|RESTRICT|NO ACTION)", blob)
        on_del = m.group(1) if m else None
    if not on_upd:
        m = re.search(r"ON UPDATE (CASCADE|SET NULL|SET DEFAULT|RESTRICT|NO ACTION)", blob)
        on_upd = m.group(1) if m else None
    return (str(on_del).upper() if on_del else None,
            str(on_upd).upper() if on_upd else None)


def _isa(node, exp, *class_names):
    """isinstance against sqlglot classes looked up by name.

    sqlglot renames and adds expression classes between releases — exp.Unique exists in
    some versions and not others — and a hard attribute reference turns a version bump
    into an AttributeError mid-parse. Looking the names up tolerantly means an unknown
    node falls through to the textual fallback instead of crashing.
    """
    cls = tuple(c for c in (getattr(exp, n, None) for n in class_names) if c is not None)
    return bool(cls) and isinstance(node, cls)


def _table_constraint(node, t, name, dialect, exp):
    """Apply one table-level constraint node to the table dict."""
    if _isa(node, exp, "PrimaryKey"):
        cols = [c.name or c.sql() for c in node.expressions]
        if cols:
            t["primary_key"] = cols
            for c in t["columns"]:
                if c["name"] in cols:
                    c["is_pk"] = True
                    c["nullable"] = False
    elif _isa(node, exp, "UniqueColumnConstraint", "Unique", "UniqueKeyProperty"):
        inner = getattr(node, "this", None)
        cols = [c.name for c in getattr(inner, "expressions", []) or []
                if hasattr(c, "name")]
        if not cols:
            cols = [c.name for c in (node.expressions or []) if hasattr(c, "name")]
        if cols and cols not in t["unique_constraints"]:
            t["unique_constraints"].append(cols)
    elif _isa(node, exp, "ForeignKey"):
        cols = [c.name for c in node.expressions]
        ref = node.args.get("reference")
        rt, rc = None, []
        if ref is not None:
            rtbl = ref.find(exp.Table)
            rt = rtbl.name if rtbl is not None else None
            sch = ref.find(exp.Schema)
            if sch is not None:
                rc = [c.name for c in sch.expressions if hasattr(c, "name")]
        on_del, on_upd = _ref_actions(node, exp)
        t["foreign_keys"].append({
            "name": name, "columns": cols, "ref_schema": None,
            "ref_table": rt or "?", "ref_columns": rc or None,
            "on_delete": on_del, "on_update": on_upd})
    elif _isa(node, exp, "Check", "CheckColumnConstraint"):
        t["checks"].append({"name": name,
                            "expression": node.sql(dialect=dialect)})
    else:
        txt = node.sql(dialect=dialect)
        # EXCLUDE ... WITH && is how PostgreSQL expresses SQL:2011 WITHOUT OVERLAPS.
        # It has to reach the model, or TEM-005 reports a missing overlap constraint on
        # a table that declares one — which is the worst kind of false positive, since
        # the reader who checks will stop trusting the rest of the report.
        if any(k in txt.upper() for k in ("CHECK", "EXCLUDE", "OVERLAP")):
            t["checks"].append({"name": name, "expression": txt})


_CREATE_RE = re.compile(
    r"create\s+table\s+(?:if\s+not\s+exists\s+)?"
    r"(?:\[?(?P<schema>[\w]+)\]?\s*\.\s*)?\[?\"?(?P<name>[\w]+)\"?\]?\s*\((?P<body>.*?)\)\s*;",
    re.I | re.S)


def _split_top(body):
    depth, cur, out = 0, [], []
    for ch in body:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == "," and depth == 0:
            out.append("".join(cur)); cur = []
        else:
            cur.append(ch)
    if cur:
        out.append("".join(cur))
    return [p.strip() for p in out if p.strip()]


def _regex_tables(sql, by_name, model):
    for m in _CREATE_RE.finditer(sql):
        t = _table(m.group("name"), m.group("schema"))
        for part in _split_top(m.group("body")):
            up = part.upper()
            if re.match(r"(CONSTRAINT\s+\w+\s+)?PRIMARY\s+KEY", up):
                cols = re.findall(r"\(([^)]*)\)", part)
                if cols:
                    t["primary_key"] = [c.strip().strip('"[]`')
                                        for c in cols[-1].split(",")]
            elif re.match(r"(CONSTRAINT\s+\w+\s+)?UNIQUE", up):
                cols = re.findall(r"\(([^)]*)\)", part)
                if cols:
                    t["unique_constraints"].append(
                        [c.strip().strip('"[]`') for c in cols[-1].split(",")])
            elif re.match(r"(CONSTRAINT\s+\w+\s+)?FOREIGN\s+KEY", up):
                fk = _parse_fk_text(part)
                if fk:
                    t["foreign_keys"].append(fk)
            elif re.match(r"(CONSTRAINT\s+\w+\s+)?(CHECK|EXCLUDE)", up):
                nm = re.search(r"constraint\s+(\w+)", part, re.I)
                t["checks"].append({"name": nm.group(1) if nm else None,
                                    "expression": part})
            else:
                cm = re.match(r'^[\["`]?(?P<col>[\w]+)[\]"`]?\s+(?P<rest>.+)$', part, re.S)
                if not cm:
                    continue
                rest = cm.group("rest")
                tm = re.match(r"^([A-Za-z][\w ]*(?:\s*\([^)]*\))?)", rest)
                raw = tm.group(1).strip() if tm else None
                length, prec, scale = _split_type(raw)
                dm = re.search(r"\bdefault\s+('[^']*'|[\w\.\-\(\)]+)", rest, re.I)
                col = _column(
                    cm.group("col"), type=raw, length=length, precision=prec,
                    scale=scale,
                    nullable=not bool(re.search(r"\bnot\s+null\b", rest, re.I)),
                    default=dm.group(1) if dm else None,
                    is_identity=bool(re.search(r"\b(identity|auto_increment|generated"
                                               r"\s+(always|by\s+default))\b", rest, re.I)),
                )
                if re.search(r"\bprimary\s+key\b", rest, re.I):
                    col["is_pk"] = True
                    col["nullable"] = False
                    t["primary_key"] = (t["primary_key"] or []) + [col["name"]]
                if re.search(r"\bunique\b", rest, re.I):
                    t["unique_constraints"].append([col["name"]])
                t["columns"].append(col)
        by_name[t["name"].lower()] = t


def _parse_fk_text(txt):
    cols = re.search(r"foreign\s+key\s*\(([^)]*)\)", txt, re.I)
    ref = re.search(r"references\s+(?:\[?(\w+)\]?\s*\.\s*)?\[?\"?(\w+)\"?\]?"
                    r"(?:\s*\(([^)]*)\))?", txt, re.I)
    if not (cols and ref):
        return None
    od = re.search(r"on\s+delete\s+(cascade|set\s+null|set\s+default|restrict|no\s+action)",
                   txt, re.I)
    ou = re.search(r"on\s+update\s+(cascade|set\s+null|set\s+default|restrict|no\s+action)",
                   txt, re.I)
    return {"name": None,
            "columns": [c.strip().strip('"[]`') for c in cols.group(1).split(",")],
            "ref_schema": ref.group(1), "ref_table": ref.group(2),
            "ref_columns": ([c.strip().strip('"[]`') for c in ref.group(3).split(",")]
                            if ref.group(3) else None),
            "on_delete": od.group(1).upper() if od else None,
            "on_update": ou.group(1).upper() if ou else None}


def _regex_alters(sql, by_name):
    pat = re.compile(r"alter\s+table\s+(?:\[?\w+\]?\s*\.\s*)?\[?\"?(\w+)\"?\]?\s+"
                     r"add\s+(.*?);", re.I | re.S)
    for m in pat.finditer(sql):
        t = by_name.get(m.group(1).lower())
        if not t:
            continue
        body, up = m.group(2), m.group(2).upper()
        if "FOREIGN KEY" in up:
            fk = _parse_fk_text(body)
            if fk:
                nm = re.search(r"constraint\s+(\w+)", body, re.I)
                fk["name"] = nm.group(1) if nm else None
                if not any(f["columns"] == fk["columns"] for f in t["foreign_keys"]):
                    t["foreign_keys"].append(fk)
        elif "PRIMARY KEY" in up:
            cols = re.findall(r"\(([^)]*)\)", body)
            if cols and not t["primary_key"]:
                t["primary_key"] = [c.strip().strip('"[]`')
                                    for c in cols[-1].split(",")]
        elif "UNIQUE" in up:
            cols = re.findall(r"\(([^)]*)\)", body)
            if cols:
                u = [c.strip().strip('"[]`') for c in cols[-1].split(",")]
                if u not in t["unique_constraints"]:
                    t["unique_constraints"].append(u)
        elif "CHECK" in up or "EXCLUDE" in up:
            nm = re.search(r"constraint\s+(\w+)", body, re.I)
            t["checks"].append({"name": nm.group(1) if nm else None, "expression": body})


def _regex_indexes(sql, by_name):
    # The trailing WHERE clause is captured because a *partial* unique index and a
    # *composite* unique index that merely includes the delete flag are opposites: the
    # first is the remedy for a soft-delete uniqueness problem, the second is the classic
    # wrong fix, since NULLs compare distinct and it permits unlimited live duplicates.
    # Without the predicate, TEM-004 cannot tell them apart and passes both.
    pat = re.compile(r"create\s+(unique\s+)?index\s+\[?\"?(\w+)\"?\]?\s+on\s+"
                     r"(?:\[?\w+\]?\s*\.\s*)?\[?\"?(\w+)\"?\]?\s*"
                     r"(?:using\s+\w+\s*)?\(([^)]*)\)"
                     r"(?:\s+where\s+(?P<pred>[^;]+))?", re.I)
    for m in pat.finditer(sql):
        t = by_name.get(m.group(3).lower())
        if not t:
            continue
        cols = [re.sub(r"\s+(asc|desc)$", "", c.strip().strip('"[]`'), flags=re.I)
                for c in m.group(4).split(",")]
        t["indexes"].append({"name": m.group(2), "columns": cols,
                             "unique": bool(m.group(1)), "kind": None,
                             "where": (m.group("pred") or "").strip() or None})


def _regex_comments(sql, by_name):
    tpat = re.compile(r"comment\s+on\s+table\s+(?:\[?\w+\]?\s*\.\s*)?\"?(\w+)\"?\s+is\s+"
                      r"'((?:[^']|'')*)'", re.I)
    for m in tpat.finditer(sql):
        t = by_name.get(m.group(1).lower())
        if t:
            t["description"] = m.group(2).replace("''", "'")
    cpat = re.compile(r"comment\s+on\s+column\s+(?:\[?\w+\]?\s*\.\s*)?\"?(\w+)\"?"
                      r"\s*\.\s*\"?(\w+)\"?\s+is\s+'((?:[^']|'')*)'", re.I)
    for m in cpat.finditer(sql):
        t = by_name.get(m.group(1).lower())
        if not t:
            continue
        for c in t["columns"]:
            if c["name"].lower() == m.group(2).lower():
                c["description"] = m.group(3).replace("''", "'")


# ======================================================================================
# information_schema CSV
# ======================================================================================

_ALIASES = {
    "table": ["table_name", "table", "entity", "entity_name", "tablename", "object_name"],
    "schema": ["table_schema", "schema", "owner", "schema_name", "module"],
    "column": ["column_name", "column", "attribute", "attribute_name", "field",
               "field_name", "columnname"],
    "type": ["data_type", "type", "datatype", "column_type", "attribute_type",
             "sql_type"],
    "nullable": ["is_nullable", "nullable", "null", "allows_null", "optional"],
    "required": ["is_required", "required", "mandatory", "not_null"],
    "length": ["character_maximum_length", "length", "max_length", "size"],
    "precision": ["numeric_precision", "precision"],
    "scale": ["numeric_scale", "scale"],
    "default": ["column_default", "default", "default_value"],
    "description": ["description", "comments", "comment", "definition", "remarks",
                    "business_definition", "meaning"],
    "pk": ["is_primary_key", "primary_key", "pk", "is_pk", "key"],
    "fk_table": ["referenced_table", "fk_table", "fk_table_name", "references_table",
                 "ref_table", "parent_table", "fktablename"],
    "fk_column": ["referenced_column", "fk_column", "fk_field_name", "ref_column",
                  "fkfieldname"],
    "classification": ["classification", "pii", "sensitivity", "data_class",
                       "confidentiality"],
}


def _resolve(headers):
    """Map canonical field -> actual header. Loud about what it could not find."""
    norm = {re.sub(r"[^a-z0-9]", "_", (h or "").strip().lower()): h for h in headers}
    out = {}
    for canon, cands in _ALIASES.items():
        for c in cands:
            if c in norm:
                out[canon] = norm[c]
                break
    return out


def parse_infoschema_csv(path):
    model = _blank_model("infoschema-csv", name=os.path.basename(path))
    model["meta"]["level"] = "physical"
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        model["meta"]["parse_warnings"].append("file contained no data rows")
        return model
    cmap = _resolve(rows[0].keys())
    for need in ("table", "column"):
        if need not in cmap:
            raise SystemExit(
                f"Could not find a '{need}' column. Looked for any of: "
                f"{', '.join(_ALIASES[need])}. Found: {', '.join(rows[0].keys())}")
    for absent in sorted(set(_ALIASES) - set(cmap)):
        model["meta"]["parse_warnings"].append(
            f"input has no '{absent}' column, so rules depending on it stand down")

    tables = {}
    for r in rows:
        tn = (r.get(cmap["table"]) or "").strip()
        if not tn:
            continue
        key = tn.lower()
        if key not in tables:
            tables[key] = _table(tn, (r.get(cmap.get("schema", "")) or None))
        t = tables[key]
        raw = (r.get(cmap.get("type", "")) or None)
        length = r.get(cmap.get("length", ""))
        prec = r.get(cmap.get("precision", ""))
        scale = r.get(cmap.get("scale", ""))
        nullable = None
        if "nullable" in cmap:
            nullable = _truthy(r.get(cmap["nullable"]))
        elif "required" in cmap:
            req = _truthy(r.get(cmap["required"]))
            nullable = (not req) if req is not None else None
        col = _column(
            (r.get(cmap["column"]) or "").strip(), type=raw,
            length=int(length) if str(length or "").strip().isdigit() else None,
            precision=int(prec) if str(prec or "").strip().isdigit() else None,
            scale=int(scale) if str(scale or "").strip().isdigit() else None,
            nullable=nullable,
            default=(r.get(cmap.get("default", "")) or None),
            description=(r.get(cmap.get("description", "")) or None),
            classification=(r.get(cmap.get("classification", "")) or None),
        )
        if _truthy(r.get(cmap.get("pk", ""))):
            col["is_pk"] = True
            t["primary_key"] = (t["primary_key"] or []) + [col["name"]]
            col["nullable"] = False
        t["columns"].append(col)
        rt = (r.get(cmap.get("fk_table", "")) or "").strip()
        if rt:
            rc = (r.get(cmap.get("fk_column", "")) or "").strip() or None
            t["foreign_keys"].append({"name": None, "columns": [col["name"]],
                                      "ref_schema": None, "ref_table": rt,
                                      "ref_columns": [rc] if rc else None,
                                      "on_delete": None, "on_update": None})
    model["tables"] = list(tables.values())
    model["meta"]["source_files"] = [os.path.abspath(path)]
    model["meta"]["parse_warnings"].append(
        "column-inventory input carries no index definitions, so index-dependent rules "
        "(PER-001, PER-002) cannot run")
    return model


# ======================================================================================
# Excel data dictionary
# ======================================================================================

def parse_excel_dd(path, sheet=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("openpyxl is required for --format excel-dd")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("the sheet is empty")

    # Find the header row rather than assuming row 1 — hand-maintained dictionaries
    # routinely carry a title block above the real header.
    header_idx, best = 0, -1
    for i, row in enumerate(rows[:20]):
        cells = [str(c).strip().lower() for c in row if c is not None]
        hits = sum(1 for c in cells
                   for cands in _ALIASES.values()
                   if re.sub(r"[^a-z0-9]", "_", c) in cands)
        if hits > best:
            header_idx, best = i, hits
    headers = [str(c) if c is not None else "" for c in rows[header_idx]]

    tmp = os.path.join(os.path.dirname(os.path.abspath(path)),
                       "._dmr_dd_tmp.csv")
    with open(tmp, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(headers)
        for row in rows[header_idx + 1:]:
            w.writerow(["" if c is None else c for c in row])
    try:
        model = parse_infoschema_csv(tmp)
    finally:
        os.remove(tmp)
    model["meta"]["source_format"] = "excel-dd"
    model["meta"]["name"] = os.path.basename(path)
    model["meta"]["source_files"] = [os.path.abspath(path)]
    if header_idx:
        model["meta"]["parse_warnings"].append(
            f"header row detected at Excel row {header_idx + 1}; rows above it were "
            f"skipped — confirm that was correct")
    return model


# ======================================================================================
# dbt YAML
# ======================================================================================

def parse_dbt_yaml(path):
    try:
        import yaml
    except ImportError:
        raise SystemExit("PyYAML is required for --format dbt-yaml")
    with open(path, encoding="utf-8") as fh:
        doc = yaml.safe_load(fh) or {}
    model = _blank_model("dbt-yaml", name=os.path.basename(path))
    model["meta"]["level"] = "logical"

    def add(node, kind):
        t = _table(node.get("name") or "?", kind=kind)
        t["description"] = node.get("description") or None
        for c in node.get("columns") or []:
            tests = [list(x)[0] if isinstance(x, dict) else x
                     for x in (c.get("tests") or c.get("data_tests") or [])]
            tests = [str(x).lower() for x in tests]
            col = _column(c.get("name") or "?",
                          type=(c.get("data_type") or None),
                          description=c.get("description") or None)
            if "not_null" in tests:
                col["nullable"] = False
            if "unique" in tests:
                t["unique_constraints"].append([col["name"]])
            for x in (c.get("tests") or c.get("data_tests") or []):
                if isinstance(x, dict) and "relationships" in x:
                    rel = x["relationships"] or {}
                    ref = str(rel.get("to") or "")
                    m = re.search(r"ref\(\s*['\"]([^'\"]+)['\"]\s*\)", ref)
                    t["foreign_keys"].append({
                        "name": None, "columns": [col["name"]], "ref_schema": None,
                        "ref_table": m.group(1) if m else ref or "?",
                        "ref_columns": [rel.get("field")] if rel.get("field") else None,
                        "on_delete": None, "on_update": None})
                if isinstance(x, dict) and "accepted_values" in x:
                    vals = (x["accepted_values"] or {}).get("values") or []
                    model["enums"].append({
                        "name": f"{t['name']}.{col['name']}",
                        "values": [str(v) for v in vals], "description": None})
            t["columns"].append(col)
        if "unique" in [str(x).lower() for x in
                        ([list(y)[0] if isinstance(y, dict) else y
                          for y in (node.get("tests") or node.get("data_tests") or [])])]:
            pass
        model["tables"].append(t)

    for m in doc.get("models") or []:
        add(m, "table")
    for s in doc.get("sources") or []:
        for tb in s.get("tables") or []:
            add(tb, "external")
    model["meta"]["source_files"] = [os.path.abspath(path)]
    model["meta"]["parse_warnings"].append(
        "a dbt schema.yml describes tests and documentation, not physical types, "
        "nullability or indexes; type, nullability and performance rules will be "
        "largely silent — that silence is missing input, not a clean model")
    return model


# ======================================================================================
# Mendix domain model (HTML documentation export)
# ======================================================================================

def parse_mendix_html(path):
    """Best-effort extraction from a Mendix domain-model documentation export.

    Mendix has no stable public text export of a domain model — the .mpr is a binary
    project file — so this reads the HTML documentation export, whose markup is not a
    contract and does change between versions. It therefore reports loudly what it did
    and did not find, and anything it could not establish stays null so the rules that
    need it stand down.

    If this adapter comes up thin against a real export, do not paper over it: build the
    model JSON by hand or from a Model Reflection extract, and use --format json. A
    hand-built model that is honest beats a scraped one that is quietly incomplete.
    """
    try:
        from html.parser import HTMLParser
    except ImportError:  # pragma: no cover
        raise SystemExit("stdlib html.parser unavailable")

    with open(path, encoding="utf-8", errors="replace") as fh:
        html = fh.read()

    class Strip(HTMLParser):
        def __init__(self):
            super().__init__()
            self.out = []

        def handle_data(self, d):
            self.out.append(d)

    model = _blank_model("mendix-html", name=os.path.basename(path))
    model["meta"]["level"] = "logical"
    model["meta"]["paradigm"] = "mendix"
    model["meta"]["source_files"] = [os.path.abspath(path)]

    s = Strip()
    s.feed(html)
    text = re.sub(r"[ \t]+", " ", "\n".join(s.out))
    lines = [ln.strip() for ln in text.splitlines()]

    entities = {}
    # Entity blocks: a heading-like line naming an entity, then attribute rows.
    ent_pat = re.compile(r"^(?:Entity|Entity:)\s*([A-Za-z][A-Za-z0-9_.]*)\s*$", re.I)
    attr_pat = re.compile(
        r"^([A-Za-z_][A-Za-z0-9_]*)\s+"
        r"(AutoNumber|Binary|Boolean|Date and time|DateTime|Decimal|Enumeration|"
        r"Hashed string|HashedString|Integer|Long|String)\b(.*)$", re.I)
    gen_pat = re.compile(r"^Generali[sz]ation:?\s*([A-Za-z][A-Za-z0-9_.]*)", re.I)
    persist_pat = re.compile(r"^Persistable:?\s*(Yes|No|True|False)", re.I)
    assoc_pat = re.compile(
        r"^([A-Za-z][A-Za-z0-9_]*_[A-Za-z][A-Za-z0-9_]*)\s+"
        r"(one-to-one|one-to-many|many-to-many|1-1|1-\*|\*-\*)\b(.*)$", re.I)
    val_pat = re.compile(r"^(Required|Unique|Equals|Range|Regular expression|"
                         r"Maximum length)\s*(?:on|:)?\s*([A-Za-z_][A-Za-z0-9_]*)?", re.I)

    cur = None
    for ln in lines:
        if not ln:
            continue
        m = ent_pat.match(ln)
        if m:
            cur = _table(m.group(1).split(".")[-1], kind="entity")
            cur["persistable"] = None
            cur["system_members"] = {"createdDate": False, "changedDate": False,
                                     "owner": False, "changedBy": False}
            cur["validation_rules"] = []
            cur["access_rules"] = []
            entities[cur["name"]] = cur
            continue
        if cur is None:
            am = assoc_pat.match(ln)
            if am:
                _add_mendix_assoc(model, am)
            continue
        m = gen_pat.match(ln)
        if m:
            cur["generalization"] = m.group(1).split(".")[-1]
            continue
        m = persist_pat.match(ln)
        if m:
            cur["persistable"] = m.group(1).lower() in {"yes", "true"}
            continue
        for sm in ("createdDate", "changedDate", "owner", "changedBy"):
            if re.match(rf"^(Store\s+)?'?{sm}'?\s*:?\s*(Yes|True)\b", ln, re.I):
                cur["system_members"][sm] = True
        m = val_pat.match(ln)
        if m and m.group(2):
            cur["validation_rules"].append({"attribute": m.group(2),
                                            "type": m.group(1), "detail": None})
            continue
        m = attr_pat.match(ln)
        if m:
            rest = m.group(3) or ""
            col = _column(m.group(1), type=m.group(2).strip(),
                          description=rest.strip() or None)
            if re.search(r"\bcalculated\b", rest, re.I):
                col["calculated"] = True
            lm = re.search(r"\b(\d+)\s*chars?\b", rest, re.I)
            if lm:
                col["length"] = int(lm.group(1))
            cur["columns"].append(col)
            continue
        am = assoc_pat.match(ln)
        if am:
            _add_mendix_assoc(model, am)

    model["tables"] = list(entities.values())
    w = model["meta"]["parse_warnings"]
    if not entities:
        w.append("no entities were recognised in this HTML export. The markup is not a "
                 "stable contract — build the model JSON by hand or from a Model "
                 "Reflection extract and use --format json.")
    else:
        w.append(f"extracted {len(entities)} entities and "
                 f"{len(model['associations'])} associations from an HTML export whose "
                 f"markup is version-dependent; spot-check against Studio Pro before "
                 f"trusting counts")
    w.append("HTML exports do not carry index definitions or access-path evidence, so "
             "MDX-005 and MDX-006 cannot run; supply used_in and indexes via --format "
             "json if index findings matter")
    if entities and not any(t["validation_rules"] for t in entities.values()):
        w.append("no validation rules were found. If the model genuinely has none that "
                 "is a serious finding (MDX-001, MDX-017); if the export omits them, "
                 "those rules are unexamined rather than failed. Establish which.")
    return model


def _add_mendix_assoc(model, m):
    mult = {"1-1": "one-to-one", "1-*": "one-to-many",
            "*-*": "many-to-many"}.get(m.group(2).lower(), m.group(2).lower())
    rest = m.group(3) or ""
    parts = m.group(1).split("_")
    dbm = re.search(r"(keep|delete associated|delete only if not associated)", rest, re.I)
    model["associations"].append({
        "name": m.group(1),
        "parent": parts[0] if parts else None,
        "child": parts[1] if len(parts) > 1 else None,
        "multiplicity": mult,
        "owner": "Both" if re.search(r"\bboth\b", rest, re.I) else "Default",
        "delete_behavior_parent": dbm.group(1) if dbm else None,
        "delete_behavior_child": None,
        "delete_behavior": dbm.group(1) if dbm else None,
        "storage": None})


# ======================================================================================
# JSON passthrough
# ======================================================================================

def parse_json(path):
    with open(path, encoding="utf-8") as fh:
        model = json.load(fh)
    model.setdefault("meta", {}).setdefault("source_format", "json")
    model["meta"].setdefault("parse_warnings", [])
    model.setdefault("tables", [])
    for t in model["tables"]:
        base = _table(t.get("name") or "?")
        base.update(t)
        t.clear()
        t.update(base)
        t["columns"] = [dict(_column(c.get("name") or "?"), **c)
                        for c in (t.get("columns") or [])]
    for k in ("enums", "associations"):
        model.setdefault(k, [])
    model.setdefault("glossary", {})
    return model


# ======================================================================================
# Driver
# ======================================================================================

ADAPTERS = {
    "ddl": lambda p, a: parse_ddl(p, a.dialect),
    "infoschema-csv": lambda p, a: parse_infoschema_csv(p),
    "excel-dd": lambda p, a: parse_excel_dd(p, a.sheet),
    "dbt-yaml": lambda p, a: parse_dbt_yaml(p),
    "mendix-html": lambda p, a: parse_mendix_html(p),
    "json": lambda p, a: parse_json(p),
}

_EXT = {".sql": "ddl", ".ddl": "ddl", ".csv": "infoschema-csv",
        ".xlsx": "excel-dd", ".xlsm": "excel-dd", ".yml": "dbt-yaml",
        ".yaml": "dbt-yaml", ".html": "mendix-html", ".htm": "mendix-html",
        ".json": "json"}


def validate(model):
    """Validate against model.schema.json if jsonschema is available."""
    try:
        import jsonschema
    except ImportError:
        return ["jsonschema not installed: output shape was not validated"]
    with open(SCHEMA_PATH, encoding="utf-8") as fh:
        schema = json.load(fh)
    v = jsonschema.Draft7Validator(schema)
    return [f"schema violation at {'/'.join(str(x) for x in e.path)}: {e.message}"
            for e in sorted(v.iter_errors(model), key=lambda e: list(e.path))][:20]


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path")
    ap.add_argument("-o", "--out", default="model.json")
    ap.add_argument("--format", choices=sorted(ADAPTERS), default=None)
    ap.add_argument("--dialect", default=None,
                    help="sqlglot dialect: oracle, postgres, tsql, mysql, snowflake ...")
    ap.add_argument("--sheet", default=None, help="worksheet name for --format excel-dd")
    ap.add_argument("--paradigm", choices=["oltp", "dim", "vault", "mendix", "doc",
                                          "graph", "semantic"])
    ap.add_argument("--level", choices=["conceptual", "logical", "physical"])
    ap.add_argument("--has-requirements", action="store_true",
                    help="a requirements document was supplied, so Correctness is "
                         "assessable")
    ap.add_argument("--has-diagram", action="store_true")
    ap.add_argument("--has-data-profile", action="store_true")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)

    fmt = args.format or _EXT.get(os.path.splitext(args.path)[1].lower())
    if not fmt:
        raise SystemExit(f"cannot infer format from '{args.path}'; pass --format")

    model = ADAPTERS[fmt](args.path, args)
    if args.paradigm:
        model["meta"]["paradigm"] = args.paradigm
    if args.level:
        model["meta"]["level"] = args.level
    model["meta"]["has_requirements_doc"] = bool(args.has_requirements)
    model["meta"]["has_diagram"] = bool(args.has_diagram)
    model["meta"]["has_data_profile"] = bool(args.has_data_profile)

    problems = validate(model)
    if problems:
        model["meta"]["parse_warnings"].extend(problems)

    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(model, fh, indent=2, ensure_ascii=False)
        fh.write("\n")

    if not args.quiet:
        print(f"format   : {fmt}")
        print(f"tables   : {len(model['tables'])}")
        print(f"columns  : {sum(len(t.get('columns') or []) for t in model['tables'])}")
        print(f"FKs      : {sum(len(t.get('foreign_keys') or []) for t in model['tables'])}")
        print(f"indexes  : {sum(len(t.get('indexes') or []) for t in model['tables'])}")
        if model["meta"]["parse_warnings"]:
            print("warnings :")
            for w in model["meta"]["parse_warnings"]:
                print(f"  - {w}")
        print(f"written  : {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
